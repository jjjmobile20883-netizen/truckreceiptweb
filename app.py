# -*- coding: utf-8 -*-
"""Toyota Tsusho Truck Receipt helper - Streamlit app."""

from io import BytesIO
from pathlib import Path
import re
from datetime import datetime
from copy import copy

import streamlit as st
import fitz
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


st.set_page_config(
    page_title="Truck Receipt Toyota Tsusho",
    page_icon="🚚",
    layout="wide",
)

st.title("🚚 Truck Receipt Toyota Tsusho")
st.caption(
    "อัปโหลดไฟล์ Excel ตัวอย่างและ PDF ระบบจะอ่านข้อมูล "
    "เติมลงใน Excel และให้ดาวน์โหลดไฟล์ใหม่"
)
st.warning(
    "โปรดตรวจสอบข้อมูลที่อ่านได้และไฟล์ Excel "
    "ก่อนนำไปใช้งานจริง โดยเฉพาะวันที่ น้ำหนัก จำนวน "
    "และเครื่องหมายการจัดส่ง"
)


def pdf_text(data: bytes) -> str:
    """อ่านข้อความจาก PDF ทุกหน้า"""
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        return "\n\n--- PDF PAGE BREAK ---\n\n".join(
            page.get_text("text", sort=True) for page in doc
        )
    finally:
        doc.close()


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def find_first(patterns, text, flags=re.I | re.M):
    for pattern in patterns:
        m = re.search(pattern, text, flags)
        if m:
            return clean(m.group(1))
    return ""


def parse_data(text: str):
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    all_text = "\n".join(lines)

    invoice = find_first(
        [
            r"Invoice\s*No\.?\s*[:#]?\s*([A-Z0-9][A-Z0-9-]{4,})",
            r"\b(THMY\d{5,})\b",
        ],
        all_text,
    )

    date = find_first(
        [
            r"Date\s*&\s*Place\s*[:]?[ \t]*(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
            r"Date\s*(?:and|&)\s*Place\s*[:]?[ \t]*"
            r"(\d{1,2}[\-/ ]+[A-Za-z0-9]{3,9}[\-/ ]+\d{4})",
        ],
        all_text,
    )

    def value_after_label(label):
        for i, line in enumerate(lines):
            m = re.search(label, line, re.I)
            if not m:
                continue

            tail = clean(line[m.end():].strip(" :\t-"))

            if tail and not re.fullmatch(r"[A-Z]{2,}", tail):
                return tail

            if i + 1 < len(lines):
                return lines[i + 1]

        return ""

    etd = value_after_label(r"^ETD\b")

    onboard = find_first(
        [
            r"^ETD\s*[:\-]?\s*(\d{1,2}\s+[A-Za-z]{3,9},?\s+\d{4})",
            r"^ONBOARD\s*[:\-]?\s*(\d{1,2}\s+[A-Za-z]{3,9},?\s+\d{4})",
        ],
        all_text,
    )

    if not onboard and re.fullmatch(
        r"\d{1,2}\s+[A-Za-z]{3,9},?\s+\d{4}",
        etd,
        re.I,
    ):
        onboard = etd

    pcs = find_first(
        [
            r"([\d,]+)\s*PC\b",
            r"\bPCS\s*[:]?[ \t]*([\d,]+)\b",
            r"([\d,]+)\s*PCS\b",
        ],
        all_text,
    )

    racks = find_first(
        [
            r"^([\d,]+)\s+Racks?\b",
            r"^RACKS?\s*[:]?[ \t]*([\d,]+)\b",
        ],
        all_text,
    )

    packages = find_first(
        [
            r"^([\d,]+)\s+PACKAGES?\b",
            r"^PACKAGES?\s*[:]?[ \t]*([\d,]+)\b",
        ],
        all_text,
    )

    net = find_first(
        [
            r"TOTAL\s+NET\s+WEIGHT\s*:?\s*([\d,]+\.\d+)",
            r"NET\s+WEIGHT\s*:?\s*([\d,]+\.\d+)",
        ],
        all_text,
    )

    gross = find_first(
        [
            r"TOTAL\s+GROSS\s+WEIGHT\s*:?\s*([\d,]+\.\d+)",
            r"GROSS\s+WEIGHT\s*:?\s*([\d,]+\.\d+)",
        ],
        all_text,
    )

    # อ่าน Shipping Marks จากหน้าที่มีหัวข้อ SHIPPING MARK ก่อน
    # เพื่อป้องกันการดึงรหัสสินค้าจากหน้าอื่นมาปะปน
    marks = []

    pages = re.split(
        r"\n\s*--- PDF PAGE BREAK ---\s*\n",
        all_text,
    )

    mark_pages = [
        page
        for page in pages
        if re.search(r"SHIPPING\s+MARK", page, re.I)
    ]

    def add_marks(source):
        for match in re.finditer(
            r"\bA0[A-Z0-9]{3,8}\b",
            source,
            re.I,
        ):
            code = match.group(0).upper()

            if code not in marks:
                marks.append(code)

    for page in mark_pages:
        marker = re.search(r"SHIPPING\s+MARK", page, re.I)

        if marker:
            add_marks(page[marker.end():marker.end() + 700])

    if not marks:
        for match in re.finditer(
            r"Case\s+No\.?\s*:?[ \t]*(A0[A-Z0-9]{3,8})",
            all_text,
            re.I,
        ):
            code = match.group(1).upper()

            if code not in marks:
                marks.append(code)

    return {
        "invoice": invoice,
        "date": date,
        "pcs": pcs.replace(",", "") if pcs else "",
        "racks": racks.replace(",", "") if racks else "",
        "packages": packages.replace(",", "") if packages else "",
        "net": net.replace(",", "") if net else "",
        "gross": gross.replace(",", "") if gross else "",
        "onboard": onboard,
        "description": "AUTOMOTIVE PARTS-DETAILS AS PER ATTACHED SHEET-",
        "marks": marks,
        "raw": all_text,
    }


def parse_date(value: str):
    if not value:
        return None

    for fmt in (
        "%d %b %Y",
        "%d %B %Y",
        "%d %b, %Y",
        "%d %B, %Y",
    ):
        try:
            return datetime.strptime(
                value.replace("  ", " "),
                fmt,
            ).date()
        except ValueError:
            pass

    return value


def writable_cell(ws, cell_ref):
    c = ws[cell_ref]

    if isinstance(c, MergedCell):
        for rng in ws.merged_cells.ranges:
            if cell_ref in rng:
                return ws.cell(rng.min_row, rng.min_col)

    return c


def set_cell(ws, ref, value):
    if value in (None, ""):
        return False

    c = writable_cell(ws, ref)
    c.value = value
    return True


def locate_label(ws, needles):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue

            if not isinstance(cell.value, str):
                continue

            txt = clean(cell.value).lower().replace(".", "")

            if any(n.lower() in txt for n in needles):
                return cell

    return None


def set_near_label(
    ws,
    needles,
    value,
    direction="right",
    max_cols=5,
):
    if value in (None, ""):
        return False

    cell = locate_label(ws, needles)

    if not cell:
        return False

    if direction == "right":
        for col in range(
            cell.column + 1,
            min(ws.max_column, cell.column + max_cols) + 1,
        ):
            target = ws.cell(cell.row, col)

            if isinstance(target, MergedCell):
                continue

            if target.value in (None, ""):
                target.value = value
                return True

        target = ws.cell(cell.row, cell.column + 1)

        if not isinstance(target, MergedCell):
            target.value = value
            return True

    return False


def format_long_date(value: str) -> str:
    d = parse_date(value)

    if hasattr(d, "strftime"):
        return d.strftime("%d %B, %Y").upper()

    return value.upper() if value else ""


def replace_date_in_label(
    text: str,
    label: str,
    date_text: str,
) -> str:
    if not text or not date_text:
        return text

    pattern = re.compile(
        r"(" + label + r"\s*:?\s*)(.*?)(?=,\s*PORT OF LOADING|$)",
        re.I,
    )

    return pattern.sub(
        lambda m: m.group(1) + date_text,
        text,
        count=1,
    )


def fill_workbook(xlsx_bytes: bytes, data: dict, chosen: str):
    wb = load_workbook(BytesIO(xlsx_bytes))
    ws = wb["TM"] if "TM" in wb.sheetnames else wb.active
    filled = []

    # ข้อมูลทั่วไปของไฟล์ CO
    if chosen == "CO" or ws.title == "AUTO":
        date_obj = parse_date(data.get("date", ""))

        if date_obj:
            if ws.title == "AUTO":
                ws["B4"] = date_obj
                ws["B4"].number_format = "dd/mm/yyyy"
                filled.append("B4 วันที่ Invoice")

            elif isinstance(ws["A4"].value, str):
                ws["A4"] = "DATE : " + date_obj.strftime("%d/%m/%Y")
                filled.append("A4 วันที่ Invoice")

        invoice = data.get("invoice", "")

        if invoice:
            for ref in ("E28", "A25", "A52"):
                current = ws[ref].value

                if isinstance(current, str):
                    if ref == "E28":
                        ws[ref] = re.sub(
                            r"(?i)(INVOICE\s*NO\.?\s*).*",
                            r"\g<1>" + invoice,
                            current,
                        )
                    else:
                        ws[ref] = f"({invoice})"

            filled.append("เลข Invoice")

        etd_long = format_long_date(data.get("onboard", ""))
        departure = ws["F11"].value

        if (
            etd_long
            and isinstance(departure, str)
            and "DEPARTURE DATE" in departure.upper()
        ):
            ws["F11"] = replace_date_in_label(
                departure,
                r"DEPARTURE\s+DATE",
                etd_long.replace(",", ""),
            )
            filled.append("F11 Departure Date จาก ETD")

        if data.get("onboard"):
            onboard_text = format_long_date(data["onboard"])
            ws["E29"] = "ONBOARD  " + onboard_text
            filled.append("E29 ONBOARD จาก ETD")

        for ref, key in (
            ("G30", "net"),
            ("G31", "gross"),
            ("F33", "pcs"),
            ("F34", "racks"),
            ("F35", "packages"),
        ):
            val = data.get(key, "")

            if val != "":
                ws[ref] = (
                    float(val)
                    if key in ("net", "gross")
                    else int(val)
                )
                filled.append(f"{ref} {key}")

        if ws.title != "AUTO":
            for ref, value in (
                ("A52", f"({invoice})" if invoice else ws["A52"].value),
                ("A53", "MALAYSIA"),
                ("A54", "MADE IN THAILAND"),
                ("A55", "REFER TO ATTACHED SHEET"),
            ):
                if value:
                    ws[ref] = value

            if ws["A52"].value:
                invoice_font = copy(ws["A52"].font)
                invoice_font.name = "Tahoma"
                invoice_font.bold = True
                ws["A52"].font = invoice_font

    # ส่วน AUTO คงการทำงานตามไฟล์ปัจจุบัน
    if ws.title == "AUTO" and chosen == "AUTO":
        invoice = data.get("invoice", "")
        date_obj = parse_date(data.get("date", ""))
        onboard_text = format_long_date(data.get("onboard", ""))

        if date_obj:
            ws["B4"] = date_obj
            ws["B4"].number_format = "dd/mm/yyyy"

        if invoice:
            ws["A25"] = f"({invoice})"
            ws["A52"] = f"({invoice})"
            ws["E28"] = f"INVOICE NO. {invoice}"

        if onboard_text:
            ws["E29"] = "ONBOARD"
            ws["F29"] = onboard_text.replace(",", "")

            for ref in ("E29", "F29"):
                c = ws[ref]
                f = copy(c.font)
                f.name = "Tahoma"
                f.sz = 11
                f.bold = True
                f.color = "000000"
                c.font = f

                a = copy(c.alignment)
                a.horizontal = "left"
                a.vertical = "center"
                a.wrap_text = False
                a.shrink_to_fit = False
                c.alignment = a

            ws.row_dimensions[29].height = max(
                ws.row_dimensions[29].height or 0,
                18,
            )

        if onboard_text and isinstance(ws["F11"].value, str):
            ws["F11"] = replace_date_in_label(
                ws["F11"].value,
                r"DEPARTURE\s+DATE",
                onboard_text.replace(",", ""),
            )

            f = copy(ws["F11"].font)
            f.bold = True
            f.color = "000000"
            ws["F11"].font = f

        for ref, key in (
            ("G30", "net"),
            ("G31", "gross"),
            ("F33", "pcs"),
            ("F34", "racks"),
            ("F35", "packages"),
        ):
            val = data.get(key, "")

            if val != "":
                ws[ref] = (
                    float(val)
                    if key in ("net", "gross")
                    else int(val)
                )

        for ref in ("G30", "G31"):
            ws[ref].number_format = "#,##0.000"

        ws["H30"] = "KGS"
        ws["H31"] = "KGS"

        ws["B22"] = "NO.OF PACKAGES"

        for ref in ("C22", "D22", "B23", "C23", "D23"):
            ws[ref] = None

        label_font = copy(ws["B22"].font)
        label_font.name = "Tahoma"
        label_font.sz = 11
        label_font.bold = True
        ws["B22"].font = label_font

        label_align = copy(ws["B22"].alignment)
        label_align.horizontal = "left"
        label_align.vertical = "center"
        label_align.wrap_text = False
        label_align.shrink_to_fit = False
        ws["B22"].alignment = label_align

    # Shipping Marks ของ AUTO เท่านั้น
    if chosen == "AUTO" and ws.title == "AUTO":
        marks = data.get("marks", [])

        for row in range(56, 128):
            for col in range(1, 6):
                cell = ws.cell(row, col)

                if isinstance(cell, MergedCell):
                    continue

                if (
                    isinstance(cell.value, str)
                    and re.fullmatch(
                        r"A0[A-Z0-9]{3,8}",
                        cell.value.strip().upper(),
                    )
                ):
                    cell.value = None

        for i, mark in enumerate(marks):
            row = 56 + i // 5
            col = 1 + i % 5
            cell = ws.cell(row, col)

            if isinstance(cell, MergedCell):
                continue

            cell.value = mark

            mark_font = copy(cell.font)
            mark_font.name = "Tahoma"
            mark_font.sz = 9
            mark_font.bold = True
            cell.font = mark_font

            mark_alignment = copy(cell.alignment)
            mark_alignment.horizontal = "left"
            cell.alignment = mark_alignment

        if marks:
            filled.append(f"Shipping Marks {len(marks)} รายการ (AUTO)")

    # Shipping Marks ของ CO เท่านั้น
    if chosen == "CO" and ws.title == "TM":
        marks = data.get("marks", [])
        start_row = 56
        start_col = 1
        columns_per_row = 5

        # ลบรหัส Shipping Marks เก่าในพื้นที่รายการ A56:E เท่านั้น
        for row in range(
            start_row,
            max(ws.max_row, start_row + 20) + 1,
        ):
            for col in range(
                start_col,
                start_col + columns_per_row,
            ):
                cell = ws.cell(row, col)

                if isinstance(cell, MergedCell):
                    continue

                if (
                    isinstance(cell.value, str)
                    and re.fullmatch(
                        r"A0[A-Z0-9]{4,8}",
                        cell.value.strip().upper(),
                    )
                ):
                    cell.value = None

        # เขียนรายการตามลำดับที่อ่านได้จาก PDF
        for i, mark in enumerate(marks):
            row = start_row + i // columns_per_row
            col = start_col + i % columns_per_row
            cell = ws.cell(row, col)

            if isinstance(cell, MergedCell):
                continue

            cell.value = mark

            mark_font = copy(cell.font)
            mark_font.name = "Tahoma"
            mark_font.bold = True
            cell.font = mark_font

            mark_alignment = copy(cell.alignment)
            mark_alignment.horizontal = "left"
            cell.alignment = mark_alignment

        if marks:
            filled.append(f"Shipping Marks {len(marks)} รายการ (CO เท่านั้น)")

    out = BytesIO()
    wb.save(out)

    return out.getvalue(), filled, ws.title


# หน้าหลักของโปรแกรม
with st.form("receipt_form"):
    chosen = st.radio(
        "เลือกรูปแบบ",
        ["AUTO", "CO"],
        horizontal=True,
        index=0,
    )

    template = st.file_uploader(
        "ไฟล์ Excel ตัวอย่าง (.xlsx)",
        type=["xlsx"],
    )

    pdf = st.file_uploader(
        "เอกสารแนบ / ใบจัดส่ง PDF",
        type=["pdf"],
    )

    submitted = st.form_submit_button(
        "อ่าน PDF และสร้างไฟล์ Excel อัตโนมัติ",
        type="primary",
    )


if submitted:
    if not template or not pdf:
        st.error("กรุณาอัปโหลดทั้งไฟล์ Excel และไฟล์ PDF ก่อนค่ะ")

    else:
        try:
            with st.spinner("กำลังอ่าน PDF และเติมข้อมูลลง Excel..."):
                text = pdf_text(pdf.getvalue())
                data = parse_data(text)

                result, filled, sheet = fill_workbook(
                    template.getvalue(),
                    data,
                    chosen,
                )

            st.subheader("ข้อมูลที่อ่านได้จาก PDF — กรุณาตรวจสอบ")

            c1, c2 = st.columns(2)

            with c1:
                st.write("**Invoice:**", data["invoice"] or "ไม่พบ")
                st.write("**วันที่ใน Invoice:**", data["date"] or "ไม่พบ")
                st.write(
                    "**วันที่ ETD/Onboard:**",
                    data["onboard"] or "ไม่พบ",
                )
                st.write("**จำนวน PCS:**", data["pcs"] or "ไม่พบ")
                st.write("**จำนวน Racks:**", data["racks"] or "ไม่พบ")

            with c2:
                st.write("**Packages:**", data["packages"] or "ไม่พบ")
                st.write("**Net weight:**", data["net"] or "ไม่พบ")
                st.write("**Gross weight:**", data["gross"] or "ไม่พบ")

                st.write(
                    f"**Shipping Marks ที่ตรวจพบ: "
                    f"{len(data['marks'])} รายการ**"
                )

                if data["marks"]:
                    for idx, mark in enumerate(data["marks"], start=1):
                        st.write(f"{idx}. `{mark}`")
                else:
                    st.error("ไม่พบ Shipping Mark ใน PDF")

                expected_marks = 24 if chosen == "CO" else 4

                if len(data["marks"]) != expected_marks:
                    st.warning(
                        f"พบ {len(data['marks'])} รายการ "
                        f"แต่รูปแบบ {chosen} นี้คาดว่าจะมี "
                        f"{expected_marks} Shipping Marks "
                        "กรุณาตรวจสอบ PDF"
                    )

            if not data["date"]:
                st.error(
                    "ยังอ่านวันที่จาก PDF ไม่ได้ "
                    "กรุณาตรวจสอบว่า PDF มีข้อความ ไม่ใช่ภาพสแกน"
                )

            if not data["marks"]:
                st.warning(
                    "ไม่พบ Shipping Mark อัตโนมัติ โปรดตรวจสอบก่อนใช้"
                )

            st.info(
                "สร้างไฟล์จากชีต: "
                + sheet
                + (
                    " | ฟิลด์ที่เติม: " + ", ".join(filled)
                    if filled
                    else " | ไม่พบตำแหน่งเซลล์ที่ตรงกับป้ายกำกับบางช่อง"
                )
            )

            stem = Path(pdf.name).stem
            output_name = f"Truck_Receipt_{chosen}_{stem}.xlsx"

            st.download_button(
                "⬇️ ดาวน์โหลด Excel ที่เติมข้อมูลแล้ว",
                data=result,
                file_name=output_name,
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                type="primary",
            )

            st.caption(
                "ระบบสร้างไฟล์ใหม่ให้ดาวน์โหลด "
                "ไม่เขียนทับไฟล์ต้นฉบับบนคอมพิวเตอร์โดยตรง "
                "เพื่อป้องกันข้อมูลต้นฉบับเสียหาย"
            )

            with st.expander(
                "ดูข้อความที่อ่านได้จาก PDF "
                "(ใช้ตรวจสอบเมื่ออ่านข้อมูลผิด)"
            ):
                st.text_area(
                    "ข้อความจาก PDF",
                    data["raw"],
                    height=300,
                )

        except Exception as e:
            st.error("เกิดข้อผิดพลาดขณะอ่านหรือสร้าง Excel")
            st.code(f"{type(e).__name__}: {e}")
            st.info(
                "ตรวจสอบว่าไฟล์ Excel เป็น .xlsx ที่เปิดได้ "
                "และไม่ถูกเปิดค้างใน Excel ระหว่างใช้งาน"
            )